from __future__ import annotations

import argparse
import base64
import binascii
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from urllib.parse import urlsplit

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageOps, UnidentifiedImageError
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent
SCHEMA = json.loads((ROOT / "resume_schema.json").read_text(encoding="utf-8"))
FONT = "微软雅黑"
BLUE = "4E7282"
BODY = "414141"
MAX_PHOTO_BYTES = 5 * 1024 * 1024
MAX_PHOTO_PIXELS = 16_000_000
MAX_TEXT_CHARACTERS = 40_000
PREVIEW_SCALE = 1400
PREVIEW_TIMEOUT = 45


class ValidationError(ValueError):
    def __init__(self, errors: dict[str, str]):
        super().__init__("请检查填写内容")
        self.errors = errors


def _photo(value: str) -> bytes:
    match = re.fullmatch(r"data:image/(jpeg|png|webp);base64,([A-Za-z0-9+/=\r\n]+)", value)
    if not match or len(match[2]) > (MAX_PHOTO_BYTES + 2) // 3 * 4 + 8:
        raise ValueError("请上传不超过 5 MB 的 JPG、PNG 或 WebP 图片")
    try:
        content = base64.b64decode(match[2], validate=True)
        if len(content) > MAX_PHOTO_BYTES:
            raise ValueError("照片超过 5 MB")
        with Image.open(io.BytesIO(content)) as image:
            if image.format not in {"JPEG", "PNG", "WEBP"}:
                raise ValueError("照片格式不受支持")
            if image.width * image.height > MAX_PHOTO_PIXELS:
                raise ValueError("请上传不超过 1600 万像素的照片")
            image.load()
            if image.format == "WEBP" or image.getexif().get(274, 1) in {2, 3, 4, 5, 6, 7, 8}:
                output = io.BytesIO()
                ImageOps.exif_transpose(image).save(output, format="PNG")
                return output.getvalue()
        return content
    except (binascii.Error, UnidentifiedImageError, OSError, Image.DecompressionBombError) as error:
        raise ValueError("无法读取照片，请重新选择有效图片") from error


def _normalize_fields(raw, fields, prefix, errors):
    if not isinstance(raw, dict):
        errors[prefix] = "请使用当前版本的字段格式"
        return {}
    known = {field["id"] for field in fields}
    if set(raw) - known:
        errors[prefix] = "发现未支持的字段，未生成简历；请检查导入文件版本"
    result = {}
    for field in fields:
        key = field["id"]
        value = raw.get(key, False if field["type"] == "checkbox" else "")
        error_key = f"{prefix}.{key}"
        if field["type"] == "checkbox":
            if not isinstance(value, bool):
                errors[error_key] = "请选择是否勾选该项"
            result[key] = value if isinstance(value, bool) else False
            continue
        if value is None:
            value = ""
        if not isinstance(value, str):
            errors[error_key] = "该项需要填写文字"
            value = ""
        value = value.strip()
        result[key] = value
        if field.get("required") and not value:
            errors[error_key] = f"请填写{field['label']}"
        if field["type"] == "file":
            if value:
                try:
                    _photo(value)
                except ValueError as error:
                    errors[error_key] = str(error)
            continue
        if len(value) > field.get("maxLength", 2000):
            errors[error_key] = f"{field['label']}最多填写 {field['maxLength']} 个字符"
        if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", value):
            errors[error_key] = "文字包含无法写入文档的控制字符，请重新粘贴"
        if not value:
            continue
        if field["type"] == "month":
            if not re.fullmatch(r"(19|20)\d{2}-(0[1-9]|1[0-2])", value):
                errors[error_key] = "请按 YYYY-MM 格式填写有效年月"
        elif field["type"] == "url":
            try:
                parsed = urlsplit(value)
                if parsed.scheme not in {"http", "https"} or not parsed.hostname or re.search(r"\s", value):
                    raise ValueError
                if parsed.username or parsed.password:
                    raise ValueError
            except ValueError:
                errors[error_key] = "请填写有效的 http:// 或 https:// 完整链接"
    return result


def normalize_payload(payload: dict) -> dict:
    errors = {}
    if not isinstance(payload, dict):
        raise ValidationError({"document": "请提供有效的简历字段对象"})
    known = {"version", "template", "language", "basic"} | {section["id"] for section in SCHEMA["sections"]}
    if set(payload) - known:
        errors["document"] = "发现未支持的模块，未生成简历；请检查导入文件版本"
    if payload.get("version", 1) != 1:
        errors["version"] = "导入文件版本不受支持"
    template = payload.get("template", "blue-gold")
    if template not in {item["id"] for item in SCHEMA["templates"]}:
        errors["template"] = "请选择蓝金或蓝灰模板"
    language = payload.get("language", "zh_CN")
    if language not in ("zh_CN", "en"):
        errors["language"] = "请选择中文版或英文版简历"
    result = {"version": 1, "template": template, "language": language}
    basic = _normalize_fields(payload.get("basic", {}), SCHEMA["basic"]["fields"], "basic", errors)
    result["basic"] = basic
    if basic.get("phone"):
        phone = basic["phone"]
        valid_characters = re.fullmatch(r"\+?[\d ().\t-]+(?:(?:x|ext\.?|分机)[ \t]*\d+)?", phone, re.IGNORECASE)
        if not valid_characters or not 5 <= len(re.findall(r"\d", phone)) <= 30:
            errors["basic.phone"] = "请填写有效的联系电话，海外号码可加区号"
    if basic.get("email") and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", basic["email"]):
        errors["basic.email"] = "请填写有效的电子邮箱地址"
    if basic.get("birth_month") and basic["birth_month"] > datetime.now().strftime("%Y-%m"):
        errors["basic.birth_month"] = "出生年月不能晚于当前月份"
    for section in SCHEMA["sections"]:
        key = section["id"]
        if not section["repeatable"]:
            result[key] = _normalize_fields(payload.get(key, {}), section["fields"], key, errors)
            continue
        entries = payload.get(key, [])
        if not isinstance(entries, list) or len(entries) > section["maxEntries"]:
            errors[key] = f"{section['label']}最多添加 {section['maxEntries']} 条"
            entries = []
        clean_entries = []
        for index, entry in enumerate(entries):
            prefix = f"{key}.{index}"
            clean = _normalize_fields(entry, section["fields"], prefix, errors)
            if not any(clean.values()):
                continue
            for field in section["fields"]:
                if field.get("requiredForEntry") and not clean.get(field["id"]):
                    errors[f"{prefix}.{field['id']}"] = f"请注明{field['label']}，或移除这条经历"
            start = clean.get("start_month", "")
            end = clean.get("end_month", "")
            if start and end and start > end:
                errors[f"{prefix}.end_month"] = "结束年月不能早于开始年月"
            if key != "education" and clean.get("ongoing") and end:
                errors[f"{prefix}.end_month"] = "勾选至今时请清空结束年月"
            clean_entries.append(clean)
        result[key] = clean_entries
    text_length = sum(len(value) for value in basic.values() if isinstance(value, str) and not value.startswith("data:image/"))
    for section in SCHEMA["sections"]:
        records = result[section["id"]] if section["repeatable"] else [result[section["id"]]]
        text_length += sum(len(value) for entry in records for value in entry.values() if isinstance(value, str))
    if text_length > MAX_TEXT_CHARACTERS:
        errors["document"] = "简历文字总量超过 40000 个字符，请精简后生成"
    if errors:
        raise ValidationError(errors)
    return result


def _style_font(style, size, color=BODY, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in list(fonts.attrib):
        if "theme" in attribute.lower():
            del fonts.attrib[attribute]
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)
    fonts.set(qn("w:hint"), "eastAsia")
    complex_size = rpr.find(qn("w:szCs"))
    if complex_size is not None:
        complex_size.set(qn("w:val"), str(round(size * 2)))


def _border(paragraph, accent, size=5):
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in {"val": "single", "sz": str(size), "space": "5", "color": accent}.items():
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    properties.append(borders)


def _table(document, rows, widths):
    table = document.add_table(rows=rows, cols=len(widths))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "nil")
        borders.append(item)
    properties.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, value in (("top", "0"), ("bottom", "0"), ("left", "0"), ("right", "70")):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:w"), value)
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    properties.append(margins)
    for column, width in zip(table.columns, widths):
        column.width = Mm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell.paragraphs[0].paragraph_format.line_spacing = Pt(16)
    return table


def _dates(entry, education=False, language="zh_CN"):
    start = entry.get("start_month", "").replace("-", ".")
    end = entry.get("end_month", "").replace("-", ".")
    if entry.get("ongoing"):
        expected, present = (" (expected)", "Present") if language == "en" else ("（预计）", "至今")
        end = end + expected if education and end else present
    if start and end:
        return f"{start} - {end}"
    return start or end


def _record(document, when, title, role):
    values = [when, title, role] if when else [title, role]
    widths = [36, 88, 50] if when else [128, 46]
    table = _table(document, 1, widths)
    for index, (cell, value) in enumerate(zip(table.rows[0].cells, values)):
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(value)
        run.bold = index == (1 if when else 0)
        if index == len(values) - 1:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _paragraphs(document, text, bullets=False, lead=""):
    lines = text.splitlines() if bullets else [text]
    for index, line in enumerate(lines):
        if not line.strip():
            continue
        paragraph = document.add_paragraph()
        if bullets:
            paragraph.paragraph_format.left_indent = Mm(4)
            paragraph.paragraph_format.first_line_indent = Mm(-3.5)
            paragraph.add_run("• ")
        if index == 0 and lead:
            paragraph.add_run(lead).bold = True
        paragraph.add_run(line)


def build_docx(payload: dict, target: Path) -> Path:
    if target.exists():
        raise FileExistsError("目标文件已存在，请使用新的输出目录")
    data = normalize_payload(payload)
    language = data["language"]
    english = language == "en"
    label_key = "labelEn" if english else "label"
    separator = ": " if english else "："
    title_text = "Resume" if english else "个人简历"
    leads = {
        "courses": "Relevant coursework: " if english else "相关课程：",
        "department": "Department / business area: " if english else "部门 / 业务方向：",
        "technology": "Technologies / methods: " if english else "技术 / 方法：",
        "link": "Link: " if english else "链接：",
        "issuer": "Institution: " if english else "机构：",
    }
    accent = next(item["accent"] for item in SCHEMA["templates"] if item["id"] == data["template"])
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(16)
    section.left_margin = section.right_margin = Mm(18)
    normal = document.styles["Normal"]
    _style_font(normal, 10.5)
    normal.paragraph_format.line_spacing = Pt(16)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True
    title_style = document.styles["Title"]
    _style_font(title_style, 24, "000000", True)
    title_style.paragraph_format.line_spacing = Pt(32)
    title_style.paragraph_format.space_after = Pt(12)
    title_style.paragraph_format.keep_with_next = True
    for child in list(title_style.element.get_or_add_pPr()):
        if child.tag == qn("w:pBdr"):
            title_style.element.get_or_add_pPr().remove(child)
    heading_style = document.styles["Heading 1"]
    _style_font(heading_style, 12, BLUE, True)
    heading_style.paragraph_format.space_before = Pt(5)
    heading_style.paragraph_format.space_after = Pt(3)
    heading_style.paragraph_format.line_spacing = Pt(17)
    heading_style.paragraph_format.keep_with_next = True
    document.core_properties.title = title_text
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    document.core_properties.created = datetime(2026, 10, 11)
    document.core_properties.modified = datetime(2026, 10, 11)
    title = document.add_paragraph(title_text, style="Title")
    title.paragraph_format.tab_stops.add_tab_stop(Mm(174), WD_TAB_ALIGNMENT.RIGHT)
    subtitle = title.add_run("\tPersonal resume")
    subtitle.font.size = Pt(11)
    subtitle.font.bold = False
    subtitle.font.color.rgb = RGBColor.from_string(BLUE)
    _border(title, accent, 12)

    def heading(label):
        paragraph = document.add_paragraph(label, style="Heading 1")
        _border(paragraph, accent)
        return paragraph

    basic = data["basic"]
    heading(SCHEMA["basic"][label_key])
    regular_fields = [field for field in SCHEMA["basic"]["fields"]
                      if field["id"] not in {"photo", "highest_education", "portfolio_links", "linkedin"} and basic.get(field["id"])]
    rows = (len(regular_fields) + 1) // 2
    has_photo = bool(basic.get("photo"))
    table = _table(document, rows, [72, 72, 30] if has_photo else [87, 87])
    for index, field in enumerate(regular_fields):
        cell = table.cell(index // 2, index % 2)
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(field[label_key] + separator)
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string("666666")
        paragraph.add_run(basic[field["id"]])
    if has_photo:
        photo_cell = table.cell(0, 2).merge(table.cell(rows - 1, 2))
        for paragraph in list(photo_cell.paragraphs)[1:]:
            photo_cell._tc.remove(paragraph._p)
        photo_cell.paragraphs[0].paragraph_format.line_spacing = 1
        photo_content = _photo(basic["photo"])
        with Image.open(io.BytesIO(photo_content)) as image:
            photo_height = min(35, max(25, rows * 16 * 25.4 / 72))
            photo_width = min(27, photo_height * image.width / image.height)
        photo_cell.paragraphs[0].add_run().add_picture(io.BytesIO(photo_content), width=Mm(photo_width))
    for key in ("highest_education", "portfolio_links", "linkedin"):
        if basic.get(key):
            label = next(field[label_key] for field in SCHEMA["basic"]["fields"] if field["id"] == key)
            _paragraphs(document, basic[key], lead=label + separator)

    for section_schema in SCHEMA["sections"]:
        key = section_schema["id"]
        entries = data[key]
        if not entries or key == "summary" and not entries.get("text"):
            continue
        heading(section_schema[label_key])
        if key == "summary":
            _paragraphs(document, entries["text"])
            continue
        for index, entry in enumerate(entries):
            if index:
                spacer = document.add_paragraph()
                spacer.paragraph_format.line_spacing = Pt(4)
                spacer.paragraph_format.space_after = Pt(0)
            if key == "education":
                _record(document, _dates(entry, True, language), entry["school"], " · ".join(filter(None, [entry["major"], entry["degree"]])))
                if entry["courses"]:
                    _paragraphs(document, entry["courses"], lead=leads["courses"])
                _paragraphs(document, entry["achievements"], bullets=True)
            elif key == "experience":
                _record(document, _dates(entry, language=language), entry["company"], entry["position"])
                if entry["department"]:
                    _paragraphs(document, entry["department"], lead=leads["department"])
                _paragraphs(document, entry["details"], bullets=True)
            elif key == "projects":
                _record(document, _dates(entry, language=language), entry["name"], entry["role"])
                if entry["technology"]:
                    _paragraphs(document, entry["technology"], lead=leads["technology"])
                if entry["link"]:
                    _paragraphs(document, entry["link"], lead=leads["link"])
                _paragraphs(document, entry["details"], bullets=True)
            elif key == "activities":
                _record(document, _dates(entry, language=language), entry["name"], entry["role"])
                _paragraphs(document, entry["details"], bullets=True)
            elif key in {"skills", "honors"}:
                _record(document, entry["date"].replace("-", "."), entry["name"], entry.get("level", ""))
                if entry["issuer"]:
                    _paragraphs(document, entry["issuer"], lead=leads["issuer"])
                _paragraphs(document, entry["description"], bullets=True)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return target


def convert_pdf(docx_path: Path, soffice_path: str | None = None) -> Path:
    executable = soffice_path or shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise RuntimeError("本地 PDF 转换程序尚未配置")
    executable = str(Path(executable).resolve())
    target = docx_path.with_suffix(".pdf")
    if target.exists():
        raise FileExistsError("目标 PDF 已存在，请使用新的输出目录")
    with tempfile.TemporaryDirectory(prefix="resume-office-") as profile:
        environment = dict(os.environ, XDG_CONFIG_HOME=profile, XDG_CACHE_HOME=profile, LANG="C.UTF-8")
        if sys.platform.startswith("linux"):
            environment["FONTCONFIG_FILE"] = str(ROOT / "fonts" / "libreoffice.conf")
            font_matcher = shutil.which("fc-match")
            if not font_matcher or not Path(environment["FONTCONFIG_FILE"]).is_file():
                raise RuntimeError("本地 PDF 中文字体配置尚未就绪")
            font_result = subprocess.run([font_matcher, "-f", "%{family}", "Noto Sans CJK SC"],
                                         env=environment, capture_output=True, timeout=10, check=False)
            if font_result.returncode or b"Noto Sans CJK SC" not in font_result.stdout:
                raise RuntimeError("本地 PDF 中文字体尚未安装，请配置 Noto Sans CJK SC")
        command = [executable, "--headless", "--nologo", "--nodefault", "--nofirststartwizard", "--norestore",
                   "-env:UserInstallation=" + Path(profile).as_uri(), "--convert-to", "pdf:writer_pdf_Export",
                   "--outdir", str(docx_path.parent.resolve()), str(docx_path.resolve())]
        try:
            result = subprocess.run(command, env=environment, capture_output=True, timeout=60, check=False)
        except subprocess.TimeoutExpired as error:
            raise RuntimeError("PDF 转换超时，请缩减内容后重试") from error
        if result.returncode or not target.is_file() or target.stat().st_size == 0:
            raise RuntimeError("PDF 转换失败，请检查本地转换程序与字体")
    if not PdfReader(target).pages:
        raise RuntimeError("PDF 没有生成有效页面")
    return target


def render_preview(pdf_path: Path, pages: int) -> list[str]:
    if pages < 1 or not pdf_path.is_file():
        raise RuntimeError("PDF 没有生成有效页面")
    if any(pdf_path.parent.glob("preview-*.png")):
        raise FileExistsError("预览文件已存在，请使用新的输出目录")
    executable = shutil.which("pdftoppm")
    if not executable:
        raise RuntimeError("本地简历预览程序尚未配置")

    with tempfile.TemporaryDirectory(prefix="resume-preview-", dir=pdf_path.parent) as temporary:
        prefix = Path(temporary) / "preview"
        command = [str(Path(executable).resolve()), "-scale-to", str(PREVIEW_SCALE), "-png",
                   str(pdf_path.resolve()), str(prefix.resolve())]
        try:
            result = subprocess.run(command, capture_output=True, timeout=PREVIEW_TIMEOUT, check=False)
        except subprocess.TimeoutExpired as error:
            raise RuntimeError("简历预览生成超时，请稍后重试") from error
        if result.returncode:
            raise RuntimeError("简历预览生成失败，请检查本地预览程序")

        images = {}
        for image_path in Path(temporary).iterdir():
            match = re.fullmatch(r"preview-(\d+)\.png", image_path.name)
            if not match or image_path.is_symlink() or not image_path.is_file():
                raise RuntimeError("简历预览文件格式不正确")
            page = int(match[1])
            if page in images:
                raise RuntimeError("简历预览页码重复")
            try:
                with Image.open(image_path) as image:
                    if image.format != "PNG" or max(image.size) > PREVIEW_SCALE:
                        raise ValueError("Invalid preview image")
                    image.verify()
                with Image.open(image_path) as image:
                    image.load()
            except (ValueError, UnidentifiedImageError, OSError, Image.DecompressionBombError) as error:
                raise RuntimeError("简历预览图片无效") from error
            images[page] = image_path

        if set(images) != set(range(1, pages + 1)):
            raise RuntimeError("简历预览页数与 PDF 不一致")
        output = [pdf_path.parent / f"preview-{page}.png" for page in range(1, pages + 1)]
        if any(path.exists() or path.is_symlink() for path in output):
            raise FileExistsError("预览文件已存在，请使用新的输出目录")
        for page, target in enumerate(output, start=1):
            os.link(images[page], target)
        return [path.name for path in output]


def main():
    parser = argparse.ArgumentParser(description="按明确字段机械生成简历 DOCX 和 PDF")
    parser.add_argument("input", type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--template", choices=[item["id"] for item in SCHEMA["templates"]])
    parser.add_argument("--language", choices=("zh_CN", "en"))
    formats = parser.add_mutually_exclusive_group()
    formats.add_argument("--docx-only", action="store_true")
    formats.add_argument("--preview", action="store_true", help="生成与 PDF 页数一致的 PNG 预览")
    parser.add_argument("--soffice")
    args = parser.parse_args()
    payload = json.loads(args.input.read_text(encoding="utf-8"))
    if args.template:
        payload["template"] = args.template
    if args.language:
        payload["language"] = args.language
    try:
        data = normalize_payload(payload)
        target = args.output_dir / ("resume-" + data["template"] + ".docx")
        build_docx(data, target)
        output = {"docx": str(target)}
        if not args.docx_only:
            pdf = convert_pdf(target, args.soffice)
            output.update(pdf=str(pdf), pages=len(PdfReader(pdf).pages))
            if args.preview:
                output["preview_files"] = render_preview(pdf, output["pages"])
        print(json.dumps(output, ensure_ascii=False))
    except ValidationError as error:
        print(json.dumps({"errors": error.errors}, ensure_ascii=False))
        raise SystemExit(2)


if __name__ == "__main__":
    main()
